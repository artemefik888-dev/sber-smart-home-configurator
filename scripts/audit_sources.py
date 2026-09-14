from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def clean(value: Any) -> Any:
    if isinstance(value, str):
        return " ".join(value.replace("\n", " / ").split())
    return value


def workbook_summary(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    cached = load_workbook(path, data_only=True, read_only=False)
    sheets: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        cached_sheet = cached[sheet.title]
        nonempty_rows: list[dict[str, Any]] = []
        for row in sheet.iter_rows():
            values = [clean(cell.value) for cell in row]
            if not any(value not in (None, "") for value in values):
                continue
            last = max(i for i, value in enumerate(values) if value not in (None, ""))
            values = values[: last + 1]
            cached_values = [clean(cached_sheet.cell(row=row[0].row, column=i + 1).value) for i in range(last + 1)]
            nonempty_rows.append(
                {
                    "row": row[0].row,
                    "values": values,
                    "cached": cached_values,
                }
            )
        sheets.append(
            {
                "title": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "merged_ranges": [str(value) for value in sheet.merged_cells.ranges],
                "image_count": len(getattr(sheet, "_images", [])),
                "nonempty_rows": nonempty_rows,
            }
        )
    return {
        "path": str(path),
        "sheets": sheets,
        "external_links": len(getattr(workbook, "_external_links", [])),
        "defined_names": [name for name in workbook.defined_names],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--sheet")
    parser.add_argument("--start", type=int, default=1)
    parser.add_argument("--end", type=int)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    report = workbook_summary(args.workbook)
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
        return

    print(f"Workbook: {report['path']}")
    print(f"External links: {report['external_links']}")
    for sheet in report["sheets"]:
        if args.sheet and sheet["title"] != args.sheet:
            continue
        print(
            f"\n[{sheet['title']}] rows={sheet['max_row']} cols={sheet['max_column']} "
            f"images={sheet['image_count']} merged={len(sheet['merged_ranges'])}"
        )
        for row in sheet["nonempty_rows"]:
            if row["row"] < args.start:
                continue
            if args.end is not None and row["row"] > args.end:
                continue
            print(f"{row['row']}: {row['values']}")
            if row["values"] != row["cached"]:
                print(f"   cached: {row['cached']}")


if __name__ == "__main__":
    main()
