from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CALCULATIONS = ROOT / "source-data" / "typical-apartments.xlsx"
CATALOG = ROOT / "source-data" / "catalog-2Q-2026.xlsx"
OUTPUT = ROOT / "app" / "data" / "default-data.json"
PRODUCT_IMAGES = ROOT / "public" / "products"

APARTMENTS = {
    "Студия 30 м²": {
        "key": "studio",
        "name": "Студия",
        "area": "30 м²",
        "layout": "studio",
    },
    "1-к квартира 40 м²": {
        "key": "one",
        "name": "Однокомнатная квартира",
        "area": "40 м²",
        "layout": "one",
    },
    "2-к квартира 60 м²": {
        "key": "two",
        "name": "Двухкомнатная квартира",
        "area": "60 м²",
        "layout": "two",
    },
    "3-к квартира 65-75 м²": {
        "key": "three",
        "name": "Трёхкомнатная квартира",
        "area": "65–75 м²",
        "layout": "three",
    },
}


def plain(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\n", " ").split())


def image_name(sku: str, extension: str) -> str:
    safe = re.sub(r"[^a-zA-Z0-9_-]", "-", sku).lower()
    return f"{safe}.{extension.lower().replace('jpeg', 'jpg')}"


def read_catalog() -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]], dict[str, Any]]:
    formulas = load_workbook(CATALOG, data_only=False, read_only=False)
    values = load_workbook(CATALOG, data_only=True, read_only=False)
    sheet = formulas["Sber"]
    value_sheet = values["Sber"]
    PRODUCT_IMAGES.mkdir(parents=True, exist_ok=True)

    image_by_row: dict[int, str] = {}
    for image in sheet._images:
        row = image.anchor._from.row + 1
        sku = plain(sheet.cell(row=row, column=2).value)
        if not sku or row in image_by_row:
            continue
        filename = image_name(sku, image.format or "png")
        (PRODUCT_IMAGES / filename).write_bytes(image._data())
        image_by_row[row] = f"/products/{filename}"

    products: list[dict[str, Any]] = []
    external_formula_count = 0
    external_formula_without_value = 0
    for row in range(2, sheet.max_row + 1):
        sku = plain(sheet.cell(row=row, column=2).value)
        if not sku:
            continue
        formula_price = sheet.cell(row=row, column=6).value
        price = value_sheet.cell(row=row, column=6).value
        if isinstance(formula_price, str) and formula_price.startswith("=") and "[" in formula_price:
            external_formula_count += 1
            if not isinstance(price, (int, float)):
                external_formula_without_value += 1
        products.append(
            {
                "sku": sku,
                "name": plain(sheet.cell(row=row, column=3).value),
                "description": plain(sheet.cell(row=row, column=5).value),
                "price": price if isinstance(price, (int, float)) else None,
                "image": image_by_row.get(row),
                "badge": plain(sheet.cell(row=row, column=7).value),
                "sourceRow": row,
            }
        )

    by_sku = {product["sku"]: product for product in products}
    counts = Counter(product["sku"] for product in products)
    diagnostics = {
        "productCount": len(products),
        "duplicateSku": sorted(sku for sku, count in counts.items() if count > 1),
        "invalidPriceSku": sorted(product["sku"] for product in products if product["price"] is None),
        "productsWithoutImage": sorted(product["sku"] for product in products if not product["image"]),
        "externalPriceFormulas": external_formula_count,
        "externalPriceFormulasWithoutValue": external_formula_without_value,
    }
    return products, by_sku, diagnostics


def detect_color(title: str) -> str | None:
    lowered = title.lower()
    if "карбон" in lowered:
        return "carbon"
    if "бел" in lowered and "цвет" in lowered:
        return "white"
    return None


def read_ready_sets(catalog_by_sku: dict[str, dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    formulas = load_workbook(CALCULATIONS, data_only=False, read_only=False)
    cached = load_workbook(CALCULATIONS, data_only=True, read_only=False)
    rows: list[dict[str, Any]] = []
    legacy_price_mismatches: list[dict[str, Any]] = []

    for sheet_name, apartment in APARTMENTS.items():
        sheet = formulas[sheet_name]
        cached_sheet = cached[sheet_name]
        color: str | None = None
        recommended = False
        for row in range(1, sheet.max_row + 1):
            first = plain(sheet.cell(row=row, column=1).value)
            sku = plain(sheet.cell(row=row, column=2).value)
            possible_color = detect_color(first) if not sku else None
            if possible_color:
                color = possible_color
                recommended = False
                continue
            if first.lower().startswith("дополнительно"):
                recommended = True
                continue
            quantity = cached_sheet.cell(row=row, column=3).value
            if not color or not sku or not isinstance(quantity, (int, float)) or quantity <= 0:
                continue

            product = catalog_by_sku.get(sku)
            old_price = cached_sheet.cell(row=row, column=4).value
            if product and isinstance(old_price, (int, float)) and old_price != product["price"]:
                legacy_price_mismatches.append(
                    {
                        "sheet": sheet_name,
                        "row": row,
                        "sku": sku,
                        "oldPrice": old_price,
                        "catalogPrice": product["price"],
                    }
                )
            rows.append(
                {
                    "id": f"{apartment['key']}-{color}-{row}",
                    "apartment": apartment["key"],
                    "color": color,
                    "sourceName": first,
                    "sku": sku,
                    "quantity": int(quantity) if float(quantity).is_integer() else quantity,
                    "room": plain(sheet.cell(row=row, column=6).value) or "Вся квартира",
                    "category": plain(sheet.cell(row=row, column=7).value) or "Без категории",
                    "gift": "акция" in plain(sheet.cell(row=row, column=8).value).lower(),
                    "recommended": recommended,
                    "sourceRow": row,
                    "valid": bool(product and isinstance(product.get("price"), (int, float))),
                }
            )

    used_skus = sorted({row["sku"] for row in rows})
    missing_skus = sorted(sku for sku in used_skus if sku not in catalog_by_sku)
    by_variant: dict[str, dict[str, Any]] = {}
    for apartment in (value["key"] for value in APARTMENTS.values()):
        for color in ("white", "carbon"):
            variant = [row for row in rows if row["apartment"] == apartment and row["color"] == color]
            by_variant[f"{apartment}:{color}"] = {
                "rows": len(variant),
                "devices": sum(row["quantity"] for row in variant if not row["recommended"]),
                "rooms": sorted({row["room"] for row in variant if not row["recommended"]}),
            }

    diagnostics = {
        "readyRowCount": len(rows),
        "usedSkuCount": len(used_skus),
        "usedSkus": used_skus,
        "missingSku": missing_skus,
        "legacyPriceMismatches": legacy_price_mismatches,
        "externalLinks": len(getattr(formulas, "_external_links", [])),
        "knownSourceIssues": [
            {
                "code": "one-room-gift-price",
                "severity": "warning",
                "message": "В белом блоке 1К для подарочного выключателя сохранена старая цена 5 990 ₽; применяется РРЦ каталога 3 990 ₽.",
            },
            {
                "code": "two-room-bath-total",
                "severity": "warning",
                "message": "Формула белого итога 2К не включает строку суммы ванной; итог пересчитывается по подробным строкам.",
            },
            {
                "code": "three-room-heading",
                "severity": "warning",
                "message": "Внутренний заголовок листа 3К ошибочно называет квартиру двухкомнатной; используется имя листа.",
            },
        ],
        "variants": by_variant,
    }
    return rows, diagnostics


def read_scenarios() -> list[dict[str, Any]]:
    workbook = load_workbook(CALCULATIONS, data_only=True, read_only=False)
    sheet = workbook["Сценарии"]
    definitions = [
        ("Свет", 5, range(7, 12)),
        ("Климат", 15, range(17, 20)),
        ("Безопасность", 23, range(25, 30)),
    ]
    result: list[dict[str, Any]] = []
    for category, intro_row, item_rows in definitions:
        result.append(
            {
                "category": category,
                "intro": plain(sheet.cell(row=intro_row, column=1).value),
                "items": [plain(sheet.cell(row=row, column=1).value) for row in item_rows if plain(sheet.cell(row=row, column=1).value)],
            }
        )
    return result


def main() -> None:
    products, catalog_by_sku, catalog_diagnostics = read_catalog()
    ready_rows, calculation_diagnostics = read_ready_sets(catalog_by_sku)
    data = {
        "meta": {
            "version": "2Q 2026",
            "priceDate": "07.04.2026",
            "currency": "RUB",
            "minimumOrder": 15000,
            "source": "Встроенные исходные данные",
        },
        "apartments": list(APARTMENTS.values()),
        "catalog": products,
        "readyRows": ready_rows,
        "scenarios": read_scenarios(),
        "diagnostics": {
            "catalog": catalog_diagnostics,
            "calculations": calculation_diagnostics,
        },
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(data["diagnostics"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
